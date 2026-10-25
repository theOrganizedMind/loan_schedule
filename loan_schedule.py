import csv
import json
import logging
import os
import tkinter as tk
from datetime import datetime, timedelta
from tkinter import filedialog, messagebox, simpledialog, ttk

import pandas as pd

try:
    import matplotlib.pyplot as plt
except ImportError:  # pragma: no cover - optional runtime dependency
    plt = None

try:
    from openpyxl import Workbook, load_workbook
except ImportError:  # pragma: no cover - optional runtime dependency
    Workbook = None
    load_workbook = None

LOGGER = logging.getLogger(__name__)
if not LOGGER.handlers:
    logging.basicConfig(level=logging.INFO)

DATE_FORMAT = "%Y-%m-%d"
TRACK_STATUSES = ("Paid", "Upcoming")
LOAN_HEADERS = [
    "Bank Name",
    "Loan Number",
    "Loan Amount",
    "Annual Interest Rate (%)",
    "Term (years)",
    "Start Date (YYYY-MM-DD)",
    "End Date (YYYY-MM-DD)",
    "Duration Remaining",
]
PAYMENT_HEADERS = [
    "Bank Name",
    "Loan Number",
    "Payment #",
    "Date",
    "Payment",
    "Principal Paid",
    "Interest Paid",
    "Remaining Balance",
    "Status",
]
REPORT_HEADERS = [
    "Bank Name",
    "Loan Number",
    "Start Date",
    "End Date",
    "Duration",
    "Status",
    "Loan Amount",
    "Remaining Balance",
]
CHART_OPTIONS = ("Overview", "By Bank", "By Loan")


def parse_date(value: str, label: str) -> datetime:
    """Parse a YYYY-MM-DD string into a datetime, raising a labeled validation error."""
    try:
        return datetime.strptime(value, DATE_FORMAT)
    except ValueError as exc:
        raise ValueError(f"{label} must use YYYY-MM-DD format.") from exc


def resolve_data_directory(base_dir: str | None = None) -> tuple[str, str]:
    """Return the preferred data directory path and visibility mode (PRIVATE or PUBLIC)."""
    root = base_dir or os.path.dirname(os.path.abspath(__file__))
    private_dir = os.path.join(root, "json_files")
    public_dir = os.path.join(root, "data")
    if os.path.isdir(private_dir):
        return private_dir, "PRIVATE"
    os.makedirs(public_dir, exist_ok=True)
    return public_dir, "PUBLIC"


def _validate_loan_inputs(principal: float, annual_rate: float, years: int, start_date: datetime) -> None:
    """Validate core loan calculation inputs before building an amortization schedule."""
    if principal <= 0:
        raise ValueError("Loan amount must be greater than 0.")
    if annual_rate < 0:
        raise ValueError("Annual interest rate cannot be negative.")
    if years <= 0:
        raise ValueError("Loan term (years) must be greater than 0.")
    if not isinstance(start_date, datetime):
        raise TypeError("Start date must be a datetime object.")


def calculate_schedule(principal: float, annual_rate: float, years: int, start_date: datetime) -> pd.DataFrame:
    """Calculate and return a payment-by-payment amortization schedule as a DataFrame."""
    _validate_loan_inputs(principal, annual_rate, years, start_date)
    monthly_rate = annual_rate / 12 / 100
    n_payments = years * 12
    if monthly_rate == 0:
        payment = principal / n_payments
    else:
        payment = principal * (monthly_rate * (1 + monthly_rate) ** n_payments) / ((1 + monthly_rate) ** n_payments - 1)

    schedule = []
    balance = principal
    payment_date = start_date
    for i in range(1, n_payments + 1):
        interest = balance * monthly_rate
        principal_paid = payment - interest
        balance -= principal_paid
        if balance < 1e-2:
            principal_paid += balance
            balance = 0
        schedule.append(
            {
                "Payment #": i,
                "Date": payment_date.strftime(DATE_FORMAT),
                "Payment": round(payment, 2),
                "Principal Paid": round(principal_paid, 2),
                "Interest Paid": round(interest, 2),
                "Remaining Balance": round(balance, 2),
            }
        )
        payment_date += timedelta(days=30)
        if balance <= 0:
            break
    return pd.DataFrame(schedule)


def format_duration_remaining(end_date_value: str, today: datetime | None = None) -> str:
    """Format remaining duration to an end date using days for short spans and years for long spans."""
    today_value = (today or datetime.today()).date()
    end_date = parse_date(end_date_value, "End Date").date()
    delta_days = (end_date - today_value).days
    if delta_days <= 0:
        return "0 days"
    if delta_days > 365:
        years = round(delta_days / 365, 1)
        return f"{years} years"
    return f"{delta_days} days"


def build_payment_rows(bank_name: str, loan_number: str, schedule: pd.DataFrame, today: datetime | None = None) -> list[dict[str, object]]:
    """Convert a schedule DataFrame into persisted payment row dictionaries with computed status."""
    today_date = (today or datetime.today()).date()
    rows: list[dict[str, object]] = []
    for _, row in schedule.iterrows():
        payment_date = parse_date(str(row["Date"]), "Payment Date").date()
        status = "Upcoming" if payment_date >= today_date else "Paid"
        rows.append(
            {
                "bank_name": bank_name,
                "loan_number": loan_number,
                "payment_number": int(row["Payment #"]),
                "date": str(row["Date"]),
                "payment": float(row["Payment"]),
                "principal_paid": float(row["Principal Paid"]),
                "interest_paid": float(row["Interest Paid"]),
                "remaining_balance": float(row["Remaining Balance"]),
                "status": status,
            }
        )
    return rows


class LoanScheduleApp:
    def __init__(self, root: tk.Tk):
        """Initialize application state, load persisted data, and build the tabbed interface."""
        self.root = root
        self.root.title("Loan Management and Tracking")
        self.root.geometry("1400x900")

        data_directory, visibility = resolve_data_directory()
        self.data_directory = data_directory
        self.data_visibility = visibility
        self.data_file = os.path.join(data_directory, f"loan_schedule_{visibility.lower()}.json")
        self.data = {"loans": [], "payments": []}
        self.current_report_rows: list[dict[str, object]] = []

        self.load_data()
        self.build_ui()
        self.refresh_all()

    def _safe_execute(self, action, *, title: str = "Operation Failed") -> None:
        """Execute a callback with centralized user-safe exception handling."""
        try:
            action()
        except ValueError as exc:
            LOGGER.warning("Validation error during %s: %s", getattr(action, "__name__", "action"), exc)
            messagebox.showwarning("Validation Error", str(exc))
        except (OSError, json.JSONDecodeError) as exc:
            LOGGER.exception("Data I/O error during %s: %s", getattr(action, "__name__", "action"), exc)
            messagebox.showerror(title, "A data access error occurred. Please try again.")
        except ImportError as exc:
            LOGGER.warning("Missing dependency during %s: %s", getattr(action, "__name__", "action"), exc)
            messagebox.showerror(title, str(exc))
        except Exception as exc:  # pragma: no cover - Tk runtime protection
            LOGGER.exception("Unexpected error during %s: %s", getattr(action, "__name__", "action"), exc)
            messagebox.showerror(title, "An unexpected application error occurred.")

    def load_data(self) -> None:
        """Load loans and payment records from the configured JSON file."""
        if not os.path.exists(self.data_file):
            self.save_data()
            return
        with open(self.data_file, "r", encoding="utf-8") as handle:
            loaded = json.load(handle)
        loans = loaded.get("loans", [])
        payments = loaded.get("payments", [])
        if not isinstance(loans, list) or not isinstance(payments, list):
            raise ValueError("Data file has invalid structure.")
        self.data = {"loans": loans, "payments": payments}

    def save_data(self) -> None:
        """Persist current in-memory loan and payment data to disk."""
        os.makedirs(self.data_directory, exist_ok=True)
        with open(self.data_file, "w", encoding="utf-8") as handle:
            json.dump(self.data, handle, indent=2)

    def build_ui(self) -> None:
        """Create the notebook container and build all tab UIs."""
        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill="both", expand=True, padx=10, pady=10)

        self.tab_loans = ttk.Frame(self.notebook)
        self.tab_track = ttk.Frame(self.notebook)
        self.tab_reports = ttk.Frame(self.notebook)
        self.tab_charts = ttk.Frame(self.notebook)

        self.notebook.add(self.tab_loans, text="📕Loan Management")
        self.notebook.add(self.tab_track, text="📍Track Loan")
        self.notebook.add(self.tab_reports, text="📄Reports")
        self.notebook.add(self.tab_charts, text="📈Charts")

        self.build_loan_tab()
        self.build_track_tab()
        self.build_reports_tab()
        self.build_charts_tab()

    def build_loan_tab(self) -> None:
        """Build the Loan Management tab fields, actions, and results table."""
        form = ttk.LabelFrame(
            self.tab_loans,
            text=f"Loan Management ({self.data_visibility} data file: {os.path.basename(self.data_file)})",
        )
        form.grid(row=0, column=0, sticky="ew", padx=10, pady=8)

        ttk.Label(form, text="Bank Name").grid(row=0, column=0, sticky="e", padx=6, pady=6)
        self.loan_bank_entry = ttk.Combobox(form, state="normal", width=24)
        self.loan_bank_entry.grid(row=0, column=1, sticky="w", padx=6, pady=6)

        ttk.Label(form, text="Loan Number").grid(row=0, column=2, sticky="e", padx=6, pady=6)
        self.loan_number_entry = ttk.Combobox(form, state="normal", width=24)
        self.loan_number_entry.grid(row=0, column=3, sticky="w", padx=6, pady=6)

        ttk.Label(form, text="Loan Amount").grid(row=0, column=4, sticky="e", padx=6, pady=6)
        self.loan_amount_entry = ttk.Entry(form, width=18)
        self.loan_amount_entry.grid(row=0, column=5, sticky="w", padx=6, pady=6)

        ttk.Label(form, text="Annual Interest Rate (%)").grid(row=1, column=0, sticky="e", padx=6, pady=6)
        self.loan_rate_entry = ttk.Entry(form, width=24)
        self.loan_rate_entry.grid(row=1, column=1, sticky="w", padx=6, pady=6)

        ttk.Label(form, text="Term (years)").grid(row=1, column=2, sticky="e", padx=6, pady=6)
        self.loan_years_entry = ttk.Entry(form, width=24)
        self.loan_years_entry.grid(row=1, column=3, sticky="w", padx=6, pady=6)

        ttk.Label(form, text="Start Date (YYYY-MM-DD)").grid(row=1, column=4, sticky="e", padx=6, pady=6)
        self.loan_start_entry = ttk.Entry(form, width=18)
        self.loan_start_entry.grid(row=1, column=5, sticky="w", padx=6, pady=6)
        self.loan_start_entry.insert(0, datetime.today().strftime(DATE_FORMAT))

        button_row = ttk.Frame(form)
        button_row.grid(row=2, column=0, columnspan=6, sticky="ew", padx=6, pady=6)

        ttk.Button(button_row, text="Add Loan", command=lambda: self._safe_execute(self.add_loan)).grid(row=0, column=0, padx=4, pady=4)
        ttk.Button(button_row, text="Update Loan", command=lambda: self._safe_execute(self.update_loan)).grid(row=0, column=1, padx=4, pady=4)
        ttk.Button(button_row, text="Remove Loan", command=lambda: self._safe_execute(self.remove_loan)).grid(row=0, column=2, padx=4, pady=4)
        ttk.Button(button_row, text="Search Loan", command=lambda: self._safe_execute(self.search_loans)).grid(row=0, column=3, padx=4, pady=4)
        ttk.Button(button_row, text="Clear Results", command=lambda: self._safe_execute(self.clear_loan_results)).grid(row=0, column=4, padx=4, pady=4)
        ttk.Button(button_row, text="Export", command=lambda: self._safe_execute(self.export_loans)).grid(row=0, column=5, padx=4, pady=4)
        button_row.columnconfigure(6, weight=1)

        table_frame = ttk.LabelFrame(self.tab_loans, text="Loan List")
        table_frame.grid(row=1, column=0, sticky="nsew", padx=10, pady=8)
        columns = ("bank_name", "loan_number", "loan_amount", "annual_rate", "term_years", "start_date", "end_date", "duration_remaining")
        self.loan_tree = ttk.Treeview(table_frame, columns=columns, show="headings", height=14)
        widths = (180, 140, 130, 160, 120, 140, 140, 160)
        for col, heading, width in zip(columns, LOAN_HEADERS, widths):
            self.loan_tree.heading(col, text=heading)
            self.loan_tree.column(col, width=width, anchor="w")
        loan_scroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.loan_tree.yview)
        self.loan_tree.configure(yscrollcommand=loan_scroll.set)
        self.loan_tree.grid(row=0, column=0, sticky="nsew", padx=6, pady=6)
        loan_scroll.grid(row=0, column=1, sticky="ns", pady=6)
        self.loan_tree.bind("<<TreeviewSelect>>", self.on_loan_selected)

        self.tab_loans.rowconfigure(1, weight=1)
        self.tab_loans.columnconfigure(0, weight=1)
        table_frame.rowconfigure(0, weight=1)
        table_frame.columnconfigure(0, weight=1)

    def build_track_tab(self) -> None:
        """Build the Track Loan tab fields, actions, and grouped payment tree."""
        form = ttk.LabelFrame(self.tab_track, text="Track Loan")
        form.grid(row=0, column=0, sticky="ew", padx=10, pady=8)

        ttk.Label(form, text="Bank Name").grid(row=0, column=0, sticky="e", padx=6, pady=6)
        self.track_bank_entry = ttk.Combobox(form, state="normal", width=24)
        self.track_bank_entry.grid(row=0, column=1, sticky="w", padx=6, pady=6)

        ttk.Label(form, text="Loan Number").grid(row=0, column=2, sticky="e", padx=6, pady=6)
        self.track_loan_entry = ttk.Combobox(form, state="normal", width=24)
        self.track_loan_entry.grid(row=0, column=3, sticky="w", padx=6, pady=6)

        ttk.Label(form, text="Status").grid(row=0, column=4, sticky="e", padx=6, pady=6)
        self.track_status_combo = ttk.Combobox(form, state="readonly", values=TRACK_STATUSES, width=18)
        self.track_status_combo.grid(row=0, column=5, sticky="w", padx=6, pady=6)
        self.track_status_combo.set("Upcoming")

        button_row = ttk.Frame(form)
        button_row.grid(row=1, column=0, columnspan=6, sticky="ew", padx=6, pady=6)
        ttk.Button(button_row, text="Mark Selected Paid", command=lambda: self._safe_execute(self.mark_selected_paid)).grid(row=0, column=0, padx=4, pady=4)
        ttk.Button(button_row, text="Refresh", command=lambda: self._safe_execute(self.refresh_track_tree)).grid(row=0, column=1, padx=4, pady=4)
        ttk.Button(button_row, text="Search", command=lambda: self._safe_execute(self.search_track_rows)).grid(row=0, column=2, padx=4, pady=4)
        ttk.Button(button_row, text="Clear", command=lambda: self._safe_execute(self.clear_track_filters)).grid(row=0, column=3, padx=4, pady=4)
        ttk.Button(button_row, text="Edit", command=lambda: self._safe_execute(self.edit_selected_payment)).grid(row=0, column=4, padx=4, pady=4)
        ttk.Button(button_row, text="Import", command=lambda: self._safe_execute(self.import_payments)).grid(row=0, column=5, padx=4, pady=4)
        ttk.Button(button_row, text="Export", command=lambda: self._safe_execute(self.export_track_rows)).grid(row=0, column=6, padx=4, pady=4)
        ttk.Button(button_row, text="Template", command=lambda: self._safe_execute(self.export_payment_template)).grid(row=0, column=7, padx=4, pady=4)
        button_row.columnconfigure(8, weight=1)

        table_frame = ttk.LabelFrame(self.tab_track, text="Upcoming Payments Grouped by Bank and Loan")
        table_frame.grid(row=1, column=0, sticky="nsew", padx=10, pady=8)
        columns = ("payment_number", "date", "payment", "principal_paid", "interest_paid", "remaining_balance")
        self.track_tree = ttk.Treeview(table_frame, columns=columns, show="tree headings", height=15)
        self.track_tree.heading("#0", text="Bank Name / Loan Number")
        self.track_tree.column("#0", width=280, anchor="w")
        headings = ("Payment #", "Date", "Payment", "Principal Paid", "Interest Paid", "Remaining Balance")
        widths = (90, 130, 120, 130, 120, 140)
        for col, heading, width in zip(columns, headings, widths):
            self.track_tree.heading(col, text=heading)
            self.track_tree.column(col, width=width, anchor="w")
        scroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.track_tree.yview)
        self.track_tree.configure(yscrollcommand=scroll.set)
        self.track_tree.grid(row=0, column=0, sticky="nsew", padx=6, pady=6)
        scroll.grid(row=0, column=1, sticky="ns", pady=6)

        self.tab_track.rowconfigure(1, weight=1)
        self.tab_track.columnconfigure(0, weight=1)
        table_frame.rowconfigure(0, weight=1)
        table_frame.columnconfigure(0, weight=1)

    def build_reports_tab(self) -> None:
        """Build the Reports tab filters, action buttons, result tree, and summary panel."""
        form = ttk.LabelFrame(self.tab_reports, text="Report Filters")
        form.grid(row=0, column=0, sticky="ew", padx=10, pady=8)

        ttk.Label(form, text="Bank Name").grid(row=0, column=0, sticky="e", padx=6, pady=6)
        self.report_bank_entry = ttk.Combobox(form, state="normal", width=30)
        self.report_bank_entry.grid(row=0, column=1, sticky="w", padx=6, pady=6)

        ttk.Label(form, text="Loan Number").grid(row=0, column=2, sticky="e", padx=6, pady=6)
        self.report_loan_entry = ttk.Combobox(form, state="normal", width=30)
        self.report_loan_entry.grid(row=0, column=3, sticky="w", padx=6, pady=6)

        button_row = ttk.Frame(form)
        button_row.grid(row=1, column=0, columnspan=8, sticky="ew", padx=6, pady=6)
        ttk.Button(button_row, text="Run Filtered Report", command=lambda: self._safe_execute(self.run_filtered_report)).grid(row=0, column=0, padx=4, pady=4)
        ttk.Button(button_row, text="All Active Loan", command=lambda: self._safe_execute(self.run_active_loan_report)).grid(row=0, column=1, padx=4, pady=4)
        ttk.Button(button_row, text="Loan History", command=lambda: self._safe_execute(self.run_loan_history_report)).grid(row=0, column=2, padx=4, pady=4)
        ttk.Button(button_row, text="Export", command=lambda: self._safe_execute(self.export_reports)).grid(row=0, column=3, padx=4, pady=4)
        ttk.Button(button_row, text="Clear", command=lambda: self._safe_execute(self.clear_report_filters)).grid(row=0, column=4, padx=4, pady=4)
        button_row.columnconfigure(5, weight=1)

        table_frame = ttk.LabelFrame(self.tab_reports, text="Report Results")
        table_frame.grid(row=1, column=0, sticky="nsew", padx=10, pady=8)
        columns = ("bank_name", "loan_number", "start_date", "end_date", "duration", "status", "loan_amount", "remaining_balance")
        self.report_tree = ttk.Treeview(table_frame, columns=columns, show="headings", height=14)
        widths = (180, 130, 120, 120, 120, 110, 130, 150)
        for col, heading, width in zip(columns, REPORT_HEADERS, widths):
            self.report_tree.heading(col, text=heading)
            self.report_tree.column(col, width=width, anchor="w")
        report_scroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.report_tree.yview)
        self.report_tree.configure(yscrollcommand=report_scroll.set)
        self.report_tree.grid(row=0, column=0, sticky="nsew", padx=6, pady=6)
        report_scroll.grid(row=0, column=1, sticky="ns", pady=6)
        self.report_tree.bind("<Double-1>", self.on_report_double_click)

        summary_frame = ttk.LabelFrame(self.tab_reports, text="Summary View")
        summary_frame.grid(row=2, column=0, sticky="nsew", padx=10, pady=(0, 8))
        self.summary_text = tk.Text(summary_frame, height=8, wrap="word", font=("Times New Roman", 12))
        self.summary_text.grid(row=0, column=0, sticky="nsew", padx=6, pady=6)
        self.summary_text.configure(state="disabled")

        self.tab_reports.rowconfigure(1, weight=1)
        self.tab_reports.rowconfigure(2, weight=1)
        self.tab_reports.columnconfigure(0, weight=1)
        table_frame.rowconfigure(0, weight=1)
        table_frame.columnconfigure(0, weight=1)
        summary_frame.rowconfigure(0, weight=1)
        summary_frame.columnconfigure(0, weight=1)

    def build_charts_tab(self) -> None:
        """Build the Charts tab filters, KPI view, and chart details table."""
        controls = ttk.LabelFrame(self.tab_charts, text="Chart Controls")
        controls.grid(row=0, column=0, sticky="ew", padx=10, pady=8)

        ttk.Label(controls, text="Bank Name").grid(row=0, column=0, sticky="e", padx=6, pady=6)
        self.chart_bank_entry = ttk.Combobox(controls, state="normal", width=22)
        self.chart_bank_entry.grid(row=0, column=1, sticky="w", padx=6, pady=6)

        ttk.Label(controls, text="Loan Number").grid(row=0, column=2, sticky="e", padx=6, pady=6)
        self.chart_loan_entry = ttk.Combobox(controls, state="normal", width=22)
        self.chart_loan_entry.grid(row=0, column=3, sticky="w", padx=6, pady=6)

        ttk.Label(controls, text="Chart Option").grid(row=0, column=4, sticky="e", padx=6, pady=6)
        self.chart_option_combo = ttk.Combobox(controls, values=CHART_OPTIONS, state="readonly", width=18)
        self.chart_option_combo.grid(row=0, column=5, sticky="w", padx=6, pady=6)
        self.chart_option_combo.set("Overview")

        ttk.Label(controls, text="Start Date").grid(row=1, column=0, sticky="e", padx=6, pady=6)
        self.chart_start_entry = ttk.Entry(controls, width=22)
        self.chart_start_entry.grid(row=1, column=1, sticky="w", padx=6, pady=6)
        self.chart_start_entry.insert(0, datetime.today().strftime(DATE_FORMAT))

        ttk.Label(controls, text="End Date").grid(row=1, column=2, sticky="e", padx=6, pady=6)
        self.chart_end_entry = ttk.Entry(controls, width=22)
        self.chart_end_entry.grid(row=1, column=3, sticky="w", padx=6, pady=6)
        self.chart_end_entry.insert(0, (datetime.today() + timedelta(days=365)).strftime(DATE_FORMAT))

        ttk.Label(controls, text="Status").grid(row=1, column=4, sticky="e", padx=6, pady=6)
        self.chart_status_combo = ttk.Combobox(controls, values=("All",) + TRACK_STATUSES, state="readonly", width=18)
        self.chart_status_combo.grid(row=1, column=5, sticky="w", padx=6, pady=6)
        self.chart_status_combo.set("All")

        button_row = ttk.Frame(controls)
        button_row.grid(row=2, column=0, columnspan=9, sticky="ew", padx=6, pady=6)
        ttk.Button(button_row, text="Search", command=lambda: self._safe_execute(self.search_chart_details)).grid(row=0, column=0, padx=4, pady=4)
        ttk.Button(button_row, text="Display Chart", command=lambda: self._safe_execute(self.display_chart)).grid(row=0, column=1, padx=4, pady=4)
        ttk.Button(button_row, text="Clear", command=lambda: self._safe_execute(self.clear_chart_filters)).grid(row=0, column=2, padx=4, pady=4)
        button_row.columnconfigure(3, weight=1)

        chart_frame = ttk.LabelFrame(self.tab_charts, text="Chart View")
        chart_frame.grid(row=1, column=0, sticky="nsew", padx=10, pady=8)
        self.chart_kpi_text = tk.Text(chart_frame, height=8, wrap="word", font=("Times New Roman", 12))
        self.chart_kpi_text.grid(row=0, column=0, sticky="nsew", padx=8, pady=8)
        self.chart_kpi_text.configure(state="disabled")

        details_frame = ttk.LabelFrame(self.tab_charts, text="Chart Details")
        details_frame.grid(row=2, column=0, sticky="nsew", padx=10, pady=(0, 8))
        columns = ("bank_name", "loan_number", "loan_amount", "remaining_balance", "remaining_duration")
        self.chart_tree = ttk.Treeview(details_frame, columns=columns, show="headings", height=8)
        headings = ("Bank Name", "Loan Number", "Loan Amount", "Remaining Balance", "Remaining Duration")
        widths = (220, 150, 140, 160, 170)
        for col, heading, width in zip(columns, headings, widths):
            self.chart_tree.heading(col, text=heading)
            self.chart_tree.column(col, width=width, anchor="w")
        chart_scroll = ttk.Scrollbar(details_frame, orient="vertical", command=self.chart_tree.yview)
        self.chart_tree.configure(yscrollcommand=chart_scroll.set)
        self.chart_tree.grid(row=0, column=0, sticky="nsew", padx=6, pady=6)
        chart_scroll.grid(row=0, column=1, sticky="ns", pady=6)

        self.tab_charts.rowconfigure(1, weight=1)
        self.tab_charts.rowconfigure(2, weight=1)
        self.tab_charts.columnconfigure(0, weight=1)
        chart_frame.rowconfigure(0, weight=1)
        chart_frame.columnconfigure(0, weight=1)
        details_frame.rowconfigure(0, weight=1)
        details_frame.columnconfigure(0, weight=1)

        self._setup_autosuggestions()

    def refresh_all(self) -> None:
        """Refresh all tab views and suggestion sources from current data."""
        self.refresh_autosuggestion_values()
        self.refresh_loan_tree(self.data["loans"])
        self.refresh_track_tree()
        self.run_loan_history_report()
        self.refresh_chart_details()

    def _setup_autosuggestions(self) -> None:
        """Register autosuggestion widgets and bind key handlers for filtering suggestions."""
        self.bank_suggestion_widgets = [
            self.loan_bank_entry,
            self.track_bank_entry,
            self.report_bank_entry,
            self.chart_bank_entry,
        ]
        self.loan_suggestion_widgets = [
            self.loan_number_entry,
            self.track_loan_entry,
            self.report_loan_entry,
            self.chart_loan_entry,
        ]
        for widget in self.bank_suggestion_widgets:
            widget.bind("<KeyRelease>", lambda _event, target=widget: self._filter_suggestions(target, "bank"))
        for widget in self.loan_suggestion_widgets:
            widget.bind("<KeyRelease>", lambda _event, target=widget: self._filter_suggestions(target, "loan"))
        self.refresh_autosuggestion_values()

    def _all_bank_names(self) -> list[str]:
        """Return a sorted unique list of bank names from loan and payment data."""
        values = {
            str(item.get("bank_name", "")).strip()
            for item in (self.data.get("loans", []) + self.data.get("payments", []))
            if str(item.get("bank_name", "")).strip()
        }
        return sorted(values, key=str.lower)

    def _all_loan_numbers(self) -> list[str]:
        """Return a sorted unique list of loan numbers from loan and payment data."""
        values = {
            str(item.get("loan_number", "")).strip()
            for item in (self.data.get("loans", []) + self.data.get("payments", []))
            if str(item.get("loan_number", "")).strip()
        }
        return sorted(values, key=str.lower)

    def refresh_autosuggestion_values(self) -> None:
        """Refresh all combobox suggestion lists with current bank and loan values."""
        bank_names = self._all_bank_names()
        loan_numbers = self._all_loan_numbers()
        if hasattr(self, "bank_suggestion_widgets"):
            for widget in self.bank_suggestion_widgets:
                widget.configure(values=bank_names)
        if hasattr(self, "loan_suggestion_widgets"):
            for widget in self.loan_suggestion_widgets:
                widget.configure(values=loan_numbers)

    def _filter_suggestions(self, widget: ttk.Combobox, kind: str) -> None:
        """Filter combobox suggestions by typed text for bank or loan values."""
        source = self._all_bank_names() if kind == "bank" else self._all_loan_numbers()
        typed = widget.get().strip().lower()
        if not typed:
            widget.configure(values=source)
            return
        filtered = [value for value in source if typed in value.lower()]
        widget.configure(values=filtered)

    def get_loan_key(self, loan: dict[str, object]) -> tuple[str, str]:
        """Return a normalized (bank_name, loan_number) key for a loan-like record."""
        return str(loan.get("bank_name", "")).strip().lower(), str(loan.get("loan_number", "")).strip().lower()

    def parse_loan_form(self) -> dict[str, object]:
        """Parse and validate loan form inputs, returning normalized fields and calculated schedule."""
        bank_name = self.loan_bank_entry.get().strip()
        loan_number = self.loan_number_entry.get().strip()
        if not bank_name:
            raise ValueError("Bank Name is required.")
        if not loan_number:
            raise ValueError("Loan Number is required.")
        principal = float(self.loan_amount_entry.get().strip())
        rate = float(self.loan_rate_entry.get().strip())
        years = int(self.loan_years_entry.get().strip())
        start_date = parse_date(self.loan_start_entry.get().strip(), "Start Date")
        schedule = calculate_schedule(principal, rate, years, start_date)
        end_date = str(schedule.iloc[-1]["Date"])
        return {
            "bank_name": bank_name,
            "loan_number": loan_number,
            "loan_amount": round(principal, 2),
            "annual_interest_rate": round(rate, 4),
            "term_years": years,
            "start_date": start_date.strftime(DATE_FORMAT),
            "end_date": end_date,
            "duration_remaining": format_duration_remaining(end_date),
            "schedule": schedule,
        }

    def add_loan(self) -> None:
        """Add a new loan and replace any existing payment rows for the same loan key."""
        parsed = self.parse_loan_form()
        new_key = (str(parsed["bank_name"]).lower(), str(parsed["loan_number"]).lower())
        for existing in self.data["loans"]:
            if self.get_loan_key(existing) == new_key:
                raise ValueError("Loan already exists. Use Update Loan instead.")
        self.data["loans"].append({k: v for k, v in parsed.items() if k != "schedule"})
        self.data["payments"] = [
            row
            for row in self.data["payments"]
            if not (
                str(row.get("bank_name", "")).lower() == new_key[0]
                and str(row.get("loan_number", "")).lower() == new_key[1]
            )
        ]
        self.data["payments"].extend(build_payment_rows(parsed["bank_name"], parsed["loan_number"], parsed["schedule"]))
        self.save_data()
        self.refresh_all()
        messagebox.showinfo("Saved", "Loan added successfully.")

    def update_loan(self) -> None:
        """Update an existing loan and regenerate its payment rows after user confirmation."""
        parsed = self.parse_loan_form()
        key = (str(parsed["bank_name"]).lower(), str(parsed["loan_number"]).lower())
        confirmed = messagebox.askokcancel(
            "Confirm Update",
            f"Update loan '{parsed['loan_number']}' for '{parsed['bank_name']}'?",
            icon="warning",
        )
        if not confirmed:
            return
        for idx, loan in enumerate(self.data["loans"]):
            if self.get_loan_key(loan) == key:
                self.data["loans"][idx] = {k: v for k, v in parsed.items() if k != "schedule"}
                break
        else:
            raise ValueError("Loan not found. Add the loan first.")
        self.data["payments"] = [
            row
            for row in self.data["payments"]
            if not (
                str(row.get("bank_name", "")).lower() == key[0]
                and str(row.get("loan_number", "")).lower() == key[1]
            )
        ]
        self.data["payments"].extend(build_payment_rows(parsed["bank_name"], parsed["loan_number"], parsed["schedule"]))
        self.save_data()
        self.refresh_all()
        messagebox.showinfo("Updated", "Loan updated successfully.")

    def remove_loan(self) -> None:
        """Remove the selected loan and associated payments after user confirmation."""
        selected = self.loan_tree.selection()
        if not selected:
            raise ValueError("Select a loan to remove.")
        values = self.loan_tree.item(selected[0], "values")
        if not values:
            raise ValueError("Select a valid loan row.")
        bank_name, loan_number = str(values[0]), str(values[1])
        confirmed = messagebox.askokcancel(
            "Confirm Remove",
            f"Remove loan '{loan_number}' for '{bank_name}'?",
            icon="warning",
        )
        if not confirmed:
            return
        key = (bank_name.lower(), loan_number.lower())
        self.data["loans"] = [loan for loan in self.data["loans"] if self.get_loan_key(loan) != key]
        self.data["payments"] = [
            row
            for row in self.data["payments"]
            if not (
                str(row.get("bank_name", "")).lower() == key[0]
                and str(row.get("loan_number", "")).lower() == key[1]
            )
        ]
        self.save_data()
        self.refresh_all()
        self.clear_loan_entries()
        messagebox.showinfo("Removed", "Loan removed successfully.")

    def search_loans(self) -> None:
        """Filter and display loan rows matching bank and loan form filters."""
        bank_filter = self.loan_bank_entry.get().strip().lower()
        loan_filter = self.loan_number_entry.get().strip().lower()
        results = []
        for loan in self.data["loans"]:
            bank_name = str(loan.get("bank_name", "")).lower()
            loan_number = str(loan.get("loan_number", "")).lower()
            if bank_filter and bank_filter not in bank_name:
                continue
            if loan_filter and loan_filter not in loan_number:
                continue
            results.append(loan)
        self.refresh_loan_tree(results)

    def clear_loan_entries(self, *, include_default_start: bool = True) -> None:
        """Clear loan form fields and optionally reset start date to today."""
        self.loan_bank_entry.delete(0, tk.END)
        self.loan_number_entry.delete(0, tk.END)
        self.loan_amount_entry.delete(0, tk.END)
        self.loan_rate_entry.delete(0, tk.END)
        self.loan_years_entry.delete(0, tk.END)
        self.loan_start_entry.delete(0, tk.END)
        if include_default_start:
            self.loan_start_entry.insert(0, datetime.today().strftime(DATE_FORMAT))

    def clear_loan_results(self) -> None:
        """Reset loan filters and restore the full loan table."""
        self.clear_loan_entries()
        self.refresh_loan_tree(self.data["loans"])

    def refresh_loan_tree(self, loans: list[dict[str, object]]) -> None:
        """Render the provided loans into the Loan Management treeview."""
        self.loan_tree.delete(*self.loan_tree.get_children())
        for loan in sorted(loans, key=lambda item: (str(item.get("bank_name", "")).lower(), str(item.get("loan_number", "")).lower())):
            end_date = str(loan.get("end_date", ""))
            self.loan_tree.insert(
                "",
                "end",
                values=(
                    loan.get("bank_name", ""),
                    loan.get("loan_number", ""),
                    f"{float(loan.get('loan_amount', 0.0)):.2f}",
                    f"{float(loan.get('annual_interest_rate', 0.0)):.4f}",
                    loan.get("term_years", ""),
                    loan.get("start_date", ""),
                    end_date,
                    format_duration_remaining(end_date) if end_date else "",
                ),
            )

    def on_loan_selected(self, _event=None) -> None:
        """Populate loan form fields from the currently selected loan row."""
        selected = self.loan_tree.selection()
        if not selected:
            return
        values = self.loan_tree.item(selected[0], "values")
        if not values:
            return
        self.clear_loan_entries(include_default_start=False)
        self.loan_bank_entry.insert(0, values[0])
        self.loan_number_entry.insert(0, values[1])
        self.loan_amount_entry.insert(0, values[2])
        self.loan_rate_entry.insert(0, values[3])
        self.loan_years_entry.insert(0, values[4])
        self.loan_start_entry.insert(0, values[5])

    def export_loans(self) -> None:
        """Export all loan rows to a user-selected CSV or XLSX file."""
        file_path = filedialog.asksaveasfilename(
            defaultextension=".csv",
            filetypes=[("CSV files", "*.csv"), ("Excel files", "*.xlsx")],
        )
        if not file_path:
            return
        rows = []
        for loan in self.data["loans"]:
            end_date = str(loan.get("end_date", ""))
            rows.append(
                [
                    loan.get("bank_name", ""),
                    loan.get("loan_number", ""),
                    loan.get("loan_amount", 0.0),
                    loan.get("annual_interest_rate", 0.0),
                    loan.get("term_years", 0),
                    loan.get("start_date", ""),
                    end_date,
                    format_duration_remaining(end_date) if end_date else "",
                ]
            )
        self._write_table(file_path, LOAN_HEADERS, rows)
        messagebox.showinfo("Export Complete", f"Loan data exported to {file_path}")

    def _filtered_future_payments(self, bank_filter: str = "", loan_filter: str = "", status_filter: str = "") -> list[dict[str, object]]:
        """Return payments filtered by bank/loan/status with future-payment rules applied."""
        today_date = datetime.today().date()
        bank_filter = bank_filter.strip().lower()
        loan_filter = loan_filter.strip().lower()
        status_filter = status_filter.strip()
        include_paid = status_filter == "Paid"
        rows = []
        for payment in self.data["payments"]:
            payment_date = parse_date(str(payment.get("date", "")), "Payment Date").date()
            payment_status = str(payment.get("status", "Upcoming"))
            if include_paid:
                if payment_status != "Paid":
                    continue
            else:
                if payment_status == "Paid":
                    continue
                if payment_date < today_date:
                    continue
            bank_name = str(payment.get("bank_name", ""))
            loan_number = str(payment.get("loan_number", ""))
            if bank_filter and bank_filter not in bank_name.lower():
                continue
            if loan_filter and loan_filter not in loan_number.lower():
                continue
            if status_filter and payment_status != status_filter:
                continue
            rows.append(payment)
        rows.sort(key=lambda item: (str(item.get("bank_name", "")).lower(), str(item.get("loan_number", "")).lower(), int(item.get("payment_number", 0))))
        return rows

    def refresh_track_tree(self) -> None:
        """Refresh the Track Loan grouped tree using current status filter selection."""
        self.track_tree.delete(*self.track_tree.get_children())
        rows = self._filtered_future_payments(status_filter=self.track_status_combo.get())
        self._insert_track_rows(rows)

    def _insert_track_rows(self, rows: list[dict[str, object]]) -> None:
        """Insert grouped payment rows and parent summary values into the Track Loan tree."""
        grouped: dict[tuple[str, str], str] = {}
        grouped_rows: dict[tuple[str, str], list[dict[str, object]]] = {}
        for payment in rows:
            key = (str(payment.get("bank_name", "")), str(payment.get("loan_number", "")))
            if key not in grouped:
                parent = self.track_tree.insert("", "end", text=f"{key[0]} | {key[1]}", values=("", "", "", "", "", ""))
                grouped[key] = parent
                grouped_rows[key] = []
            grouped_rows[key].append(payment)
            item_id = f"{key[0]}::{key[1]}::{int(payment.get('payment_number', 0))}"
            self.track_tree.insert(
                grouped[key],
                "end",
                iid=item_id,
                text="",
                values=(
                    int(payment.get("payment_number", 0)),
                    payment.get("date", ""),
                    f"{float(payment.get('payment', 0.0)):.2f}",
                    f"{float(payment.get('principal_paid', 0.0)):.2f}",
                    f"{float(payment.get('interest_paid', 0.0)):.2f}",
                    f"{float(payment.get('remaining_balance', 0.0)):.2f}",
                ),
            )
        for key, parent in grouped.items():
            group_items = grouped_rows[key]
            if not group_items:
                continue
            avg_payment = sum(float(item.get("payment", 0.0)) for item in group_items) / len(group_items)
            avg_principal = sum(float(item.get("principal_paid", 0.0)) for item in group_items) / len(group_items)
            avg_interest = sum(float(item.get("interest_paid", 0.0)) for item in group_items) / len(group_items)
            loan_info = next(
                (
                    loan
                    for loan in self.data["loans"]
                    if str(loan.get("bank_name", "")).lower() == key[0].lower()
                    and str(loan.get("loan_number", "")).lower() == key[1].lower()
                ),
                None,
            )
            loan_amount = float(loan_info.get("loan_amount", 0.0)) if loan_info else 0.0
            _status, total_remaining = self._loan_status_and_remaining(key[0], key[1], loan_amount)
            self.track_tree.item(
                parent,
                values=(
                    "",
                    "",
                    f"{avg_payment:.2f}",
                    f"{avg_principal:.2f}",
                    f"{avg_interest:.2f}",
                    f"{total_remaining:.2f}",
                ),
            )
        for parent in grouped.values():
            self.track_tree.item(parent, open=True)

    def search_track_rows(self) -> None:
        """Apply Track Loan filters and render matching grouped payment rows."""
        rows = self._filtered_future_payments(
            bank_filter=self.track_bank_entry.get(),
            loan_filter=self.track_loan_entry.get(),
            status_filter=self.track_status_combo.get(),
        )
        self.track_tree.delete(*self.track_tree.get_children())
        self._insert_track_rows(rows)

    def clear_track_filters(self) -> None:
        """Clear Track Loan filters and restore default upcoming payment view."""
        self.track_bank_entry.delete(0, tk.END)
        self.track_loan_entry.delete(0, tk.END)
        self.track_status_combo.set("Upcoming")
        self.refresh_track_tree()

    def _selected_payment_record(self) -> dict[str, object]:
        """Return the payment record represented by the currently selected tree item."""
        selected = self.track_tree.selection()
        if not selected:
            raise ValueError("Select a payment row.")
        item_id = selected[0]
        if "::" not in item_id:
            raise ValueError("Select a payment entry, not a group header.")
        bank_name, loan_number, payment_number = item_id.split("::", maxsplit=2)
        payment_number_value = int(payment_number)
        for payment in self.data["payments"]:
            if (
                str(payment.get("bank_name", "")) == bank_name
                and str(payment.get("loan_number", "")) == loan_number
                and int(payment.get("payment_number", 0)) == payment_number_value
            ):
                return payment
        raise ValueError("Selected payment record no longer exists.")

    def mark_selected_paid(self) -> None:
        """Mark the selected payment as paid and refresh dependent views."""
        payment = self._selected_payment_record()
        payment["status"] = "Paid"
        self.save_data()
        self.refresh_autosuggestion_values()
        self.refresh_track_tree()
        self.run_filtered_report()
        self.refresh_chart_details()
        messagebox.showinfo("Updated", "Selected payment marked as Paid.")

    def edit_selected_payment(self) -> None:
        """Edit selected payment date and status using dialog prompts."""
        payment = self._selected_payment_record()
        new_date = simpledialog.askstring("Edit Payment", "Date (YYYY-MM-DD):", initialvalue=str(payment.get("date", "")))
        if new_date is None:
            return
        parse_date(new_date, "Payment Date")
        new_status = simpledialog.askstring("Edit Payment", "Status (Paid/Upcoming):", initialvalue=str(payment.get("status", "Upcoming")))
        if new_status is None:
            return
        if new_status not in TRACK_STATUSES:
            raise ValueError("Status must be 'Paid' or 'Upcoming'.")
        payment["date"] = new_date
        payment["status"] = new_status
        self.save_data()
        self.refresh_autosuggestion_values()
        self.refresh_track_tree()
        self.run_filtered_report()
        self.refresh_chart_details()

    def import_payments(self) -> None:
        """Import payment rows from CSV/XLSX and merge them into existing payment data."""
        file_path = filedialog.askopenfilename(filetypes=[("CSV files", "*.csv"), ("Excel files", "*.xlsx")])
        if not file_path:
            return
        rows = self._read_table(file_path)
        if not rows:
            raise ValueError("Selected file has no rows.")
        required = set(PAYMENT_HEADERS)
        for row in rows:
            if not required.issubset(set(row.keys())):
                raise ValueError("Import file is missing required payment columns.")
        for row in rows:
            status = str(row.get("Status", "Upcoming")).strip() or "Upcoming"
            if status not in TRACK_STATUSES:
                raise ValueError("Payment status must be Paid or Upcoming.")
            record = {
                "bank_name": str(row.get("Bank Name", "")).strip(),
                "loan_number": str(row.get("Loan Number", "")).strip(),
                "payment_number": int(float(row.get("Payment #", 0))),
                "date": str(row.get("Date", "")).strip(),
                "payment": float(row.get("Payment", 0.0)),
                "principal_paid": float(row.get("Principal Paid", 0.0)),
                "interest_paid": float(row.get("Interest Paid", 0.0)),
                "remaining_balance": float(row.get("Remaining Balance", 0.0)),
                "status": status,
            }
            parse_date(record["date"], "Payment Date")
            existing = None
            for payment in self.data["payments"]:
                if (
                    str(payment.get("bank_name", "")) == record["bank_name"]
                    and str(payment.get("loan_number", "")) == record["loan_number"]
                    and int(payment.get("payment_number", 0)) == record["payment_number"]
                ):
                    existing = payment
                    break
            if existing is None:
                self.data["payments"].append(record)
            else:
                existing.update(record)
        self.save_data()
        self.refresh_autosuggestion_values()
        self.refresh_track_tree()
        self.run_filtered_report()
        self.refresh_chart_details()
        messagebox.showinfo("Import Complete", f"Imported payment rows from {file_path}")

    def export_track_rows(self) -> None:
        """Export currently filtered Track Loan rows to CSV or XLSX."""
        file_path = filedialog.asksaveasfilename(
            defaultextension=".csv",
            filetypes=[("CSV files", "*.csv"), ("Excel files", "*.xlsx")],
        )
        if not file_path:
            return
        rows = self._filtered_future_payments(
            bank_filter=self.track_bank_entry.get(),
            loan_filter=self.track_loan_entry.get(),
            status_filter=self.track_status_combo.get(),
        )
        serialized = [
            [
                row.get("bank_name", ""),
                row.get("loan_number", ""),
                row.get("payment_number", 0),
                row.get("date", ""),
                row.get("payment", 0.0),
                row.get("principal_paid", 0.0),
                row.get("interest_paid", 0.0),
                row.get("remaining_balance", 0.0),
                row.get("status", "Upcoming"),
            ]
            for row in rows
        ]
        self._write_table(file_path, PAYMENT_HEADERS, serialized)
        messagebox.showinfo("Export Complete", f"Track rows exported to {file_path}")

    def export_payment_template(self) -> None:
        """Export an empty payment template file with required headers."""
        file_path = filedialog.asksaveasfilename(
            defaultextension=".csv",
            filetypes=[("CSV files", "*.csv"), ("Excel files", "*.xlsx")],
        )
        if not file_path:
            return
        self._write_table(file_path, PAYMENT_HEADERS, [])
        messagebox.showinfo("Template Saved", f"Template saved to {file_path}")

    def _loan_status_and_remaining(self, bank_name: str, loan_number: str, loan_amount: float) -> tuple[str, float]:
        """Compute aggregate loan status and current remaining balance for a loan key."""
        today_date = datetime.today().date()
        loan_payments = [
            payment
            for payment in self.data["payments"]
            if str(payment.get("bank_name", "")).lower() == bank_name.lower()
            and str(payment.get("loan_number", "")).lower() == loan_number.lower()
        ]
        if not loan_payments:
            return "Upcoming", loan_amount
        loan_payments.sort(key=lambda item: int(item.get("payment_number", 0)))
        upcoming = [
            item
            for item in loan_payments
            if str(item.get("status", "Upcoming")) == "Upcoming"
            and parse_date(str(item.get("date", "")), "Payment Date").date() >= today_date
        ]
        if not upcoming:
            return "Paid", 0.0
        next_payment = upcoming[0]
        next_number = int(next_payment.get("payment_number", 0))
        if next_number <= 1:
            return "Upcoming", loan_amount
        previous = None
        for row in loan_payments:
            if int(row.get("payment_number", 0)) == next_number - 1:
                previous = row
                break
        if previous is None:
            return "Upcoming", loan_amount
        return "Upcoming", float(previous.get("remaining_balance", loan_amount))

    def build_report_rows(self, only_active: bool = False, apply_filter: bool = True) -> list[dict[str, object]]:
        """Build normalized report rows from loans with optional active-only and form filtering."""
        bank_filter = self.report_bank_entry.get().strip().lower() if apply_filter else ""
        loan_filter = self.report_loan_entry.get().strip().lower() if apply_filter else ""
        rows = []
        for loan in self.data["loans"]:
            bank_name = str(loan.get("bank_name", ""))
            loan_number = str(loan.get("loan_number", ""))
            if bank_filter and bank_filter not in bank_name.lower():
                continue
            if loan_filter and loan_filter not in loan_number.lower():
                continue
            status, remaining = self._loan_status_and_remaining(bank_name, loan_number, float(loan.get("loan_amount", 0.0)))
            if only_active and status != "Upcoming":
                continue
            start_date = str(loan.get("start_date", ""))
            end_date = str(loan.get("end_date", ""))
            duration_days = max((parse_date(end_date, "End Date") - parse_date(start_date, "Start Date")).days, 0)
            duration = f"{round(duration_days / 365, 1)} years" if duration_days > 365 else f"{duration_days} days"
            rows.append(
                {
                    "bank_name": bank_name,
                    "loan_number": loan_number,
                    "start_date": start_date,
                    "end_date": end_date,
                    "duration": duration,
                    "status": status,
                    "loan_amount": float(loan.get("loan_amount", 0.0)),
                    "remaining_balance": round(remaining, 2),
                }
            )
        rows.sort(key=lambda item: (item["bank_name"].lower(), item["loan_number"].lower()))
        return rows

    def _render_report_rows(self, rows: list[dict[str, object]], mode: str) -> None:
        """Render report rows into the tree and update the summary view text."""
        self.current_report_rows = rows
        self.report_tree.delete(*self.report_tree.get_children())
        for row in rows:
            self.report_tree.insert(
                "",
                "end",
                values=(
                    row["bank_name"],
                    row["loan_number"],
                    row["start_date"],
                    row["end_date"],
                    row["duration"],
                    row["status"],
                    f"{row['loan_amount']:.2f}",
                    f"{row['remaining_balance']:.2f}",
                ),
            )
        total_open = sum(1 for row in rows if row["status"] == "Upcoming")
        total_loan_amount = sum(float(row["loan_amount"]) for row in rows)
        total_remaining = sum(float(row["remaining_balance"]) for row in rows)
        row_keys = {(str(row["bank_name"]).lower(), str(row["loan_number"]).lower()) for row in rows}
        paid_count = 0
        upcoming_count = 0
        for payment in self.data["payments"]:
            key = (str(payment.get("bank_name", "")).lower(), str(payment.get("loan_number", "")).lower())
            if key not in row_keys:
                continue
            if str(payment.get("status", "Upcoming")) == "Paid":
                paid_count += 1
            else:
                upcoming_count += 1
        summary = (
            f"Data Mode: {self.data_visibility}\n"
            f"Total Open Loans: {total_open}\n"
            f"Total Loan Amount: ${total_loan_amount:,.2f}\n"
            f"Remaining Loan Balance: ${total_remaining:,.2f}\n"
            f"Payment Status Counts:\n"
            f"\tPaid: {paid_count},\n"
            f"\tUpcoming: {upcoming_count}\n"
            f"Rows in view: {len(rows)} ({mode})"
        )
        self.summary_text.configure(state="normal")
        self.summary_text.delete("1.0", tk.END)
        self.summary_text.insert(tk.END, summary)
        self.summary_text.configure(state="disabled")

    def run_filtered_report(self) -> None:
        """Run and render the reports tab using current bank/loan filters."""
        rows = self.build_report_rows(only_active=False, apply_filter=True)
        self._render_report_rows(rows, "Filtered Report")

    def clear_report_filters(self) -> None:
        """Clear report filter entries and restore full loan history report view."""
        self.report_bank_entry.delete(0, tk.END)
        self.report_loan_entry.delete(0, tk.END)
        self.run_loan_history_report()

    def on_report_double_click(self, _event=None) -> None:
        """Populate report filters from a double-clicked report row."""
        selected = self.report_tree.selection()
        if not selected:
            return
        values = self.report_tree.item(selected[0], "values")
        if not values:
            return
        self.report_bank_entry.delete(0, tk.END)
        self.report_bank_entry.insert(0, str(values[0]))
        self.report_loan_entry.delete(0, tk.END)
        self.report_loan_entry.insert(0, str(values[1]))

    def run_active_loan_report(self) -> None:
        """Run and render report rows limited to active (upcoming) loans."""
        rows = self.build_report_rows(only_active=True, apply_filter=False)
        self._render_report_rows(rows, "All Active Loan")

    def run_loan_history_report(self) -> None:
        """Run and render report rows for complete loan history."""
        rows = self.build_report_rows(only_active=False, apply_filter=False)
        self._render_report_rows(rows, "Loan History")

    def export_reports(self) -> None:
        """Export the currently rendered report rows to CSV or XLSX."""
        if not self.current_report_rows:
            raise ValueError("Run a report before exporting.")
        file_path = filedialog.asksaveasfilename(
            defaultextension=".csv",
            filetypes=[("CSV files", "*.csv"), ("Excel files", "*.xlsx")],
        )
        if not file_path:
            return
        rows = [
            [
                row["bank_name"],
                row["loan_number"],
                row["start_date"],
                row["end_date"],
                row["duration"],
                row["status"],
                row["loan_amount"],
                row["remaining_balance"],
            ]
            for row in self.current_report_rows
        ]
        self._write_table(file_path, REPORT_HEADERS, rows)
        messagebox.showinfo("Export Complete", f"Report exported to {file_path}")

    def _chart_filtered_rows(self) -> list[dict[str, object]]:
        """Return report rows filtered for charting by date range, bank/loan, and status."""
        start_date = parse_date(self.chart_start_entry.get().strip(), "Chart Start Date").date()
        end_date = parse_date(self.chart_end_entry.get().strip(), "Chart End Date").date()
        if end_date < start_date:
            raise ValueError("Chart end date must be after or equal to chart start date.")
        bank_filter = self.chart_bank_entry.get().strip().lower()
        loan_filter = self.chart_loan_entry.get().strip().lower()
        status_filter = self.chart_status_combo.get().strip()
        rows = self.build_report_rows(only_active=False, apply_filter=False)
        filtered = []
        for row in rows:
            row_start = parse_date(str(row["start_date"]), "Start Date").date()
            row_end = parse_date(str(row["end_date"]), "End Date").date()
            if row_end < start_date or row_start > end_date:
                continue
            if bank_filter and bank_filter != str(row["bank_name"]).strip().lower():
                continue
            if loan_filter and loan_filter != str(row["loan_number"]).strip().lower():
                continue
            if status_filter != "All" and row["status"] != status_filter:
                continue
            filtered.append(row)
        return filtered

    def search_chart_details(self) -> None:
        """Apply bank/loan chart search, auto-set date bounds, and refresh chart views."""
        bank_filter = self.chart_bank_entry.get().strip().lower()
        loan_filter = self.chart_loan_entry.get().strip().lower()
        matching_loans = [
            loan
            for loan in self.data["loans"]
            if (not bank_filter or bank_filter == str(loan.get("bank_name", "")).strip().lower())
            and (not loan_filter or loan_filter == str(loan.get("loan_number", "")).strip().lower())
        ]
        if matching_loans:
            start_dates = [parse_date(str(loan.get("start_date", "")), "Start Date") for loan in matching_loans]
            end_dates = [parse_date(str(loan.get("end_date", "")), "End Date") for loan in matching_loans]
            self.chart_start_entry.delete(0, tk.END)
            self.chart_start_entry.insert(0, min(start_dates).strftime(DATE_FORMAT))
            self.chart_end_entry.delete(0, tk.END)
            self.chart_end_entry.insert(0, max(end_dates).strftime(DATE_FORMAT))
        self.refresh_chart_details()

    def refresh_chart_details(self) -> None:
        """Refresh chart details table and KPI text using current chart filters."""
        rows = self._chart_filtered_rows()
        self.chart_tree.delete(*self.chart_tree.get_children())
        for row in rows:
            self.chart_tree.insert(
                "",
                "end",
                values=(
                    row["bank_name"],
                    row["loan_number"],
                    f"{float(row['loan_amount']):.2f}",
                    f"{float(row['remaining_balance']):.2f}",
                    format_duration_remaining(str(row["end_date"])),
                ),
            )
        self._update_chart_kpi(rows)

    def _update_chart_kpi(self, rows: list[dict[str, object]]) -> None:
        """Update chart KPI summary text from the supplied chart-filtered rows."""
        total_loan_amount = sum(float(row["loan_amount"]) for row in rows)
        total_remaining = sum(float(row["remaining_balance"]) for row in rows)
        avg_loan_amount = (total_loan_amount / len(rows)) if rows else 0.0
        avg_remaining = (total_remaining / len(rows)) if rows else 0.0
        avg_interest_rate = 0.0
        avg_payment_amount = 0.0
        if rows:
            key_set = {(str(row["bank_name"]).lower(), str(row["loan_number"]).lower()) for row in rows}
            matching = [
                loan
                for loan in self.data["loans"]
                if (str(loan.get("bank_name", "")).lower(), str(loan.get("loan_number", "")).lower()) in key_set
            ]
            if matching:
                avg_interest_rate = sum(float(loan.get("annual_interest_rate", 0.0)) for loan in matching) / len(matching)
            first_payments = []
            for key in key_set:
                key_rows = [
                    payment
                    for payment in self.data["payments"]
                    if str(payment.get("bank_name", "")).lower() == key[0]
                    and str(payment.get("loan_number", "")).lower() == key[1]
                ]
                if not key_rows:
                    continue
                key_rows.sort(key=lambda item: int(item.get("payment_number", 0)))
                first_payments.append(float(key_rows[0].get("payment", 0.0)))
            if first_payments:
                avg_payment_amount = sum(first_payments) / len(first_payments)
        kpi = (
            f"Number of Loans: {len(rows)}\n"
            f"Total Loan Amount: ${total_loan_amount:,.2f}\n"
            f"Total Remaining Balance: ${total_remaining:,.2f}\n"
            f"Average Loan Amount: ${avg_loan_amount:,.2f}\n"
            f"Average Payment Amount: ${avg_payment_amount:,.2f}\n"
            f"Average Remaining Balance: ${avg_remaining:,.2f}\n"
            f"Average Interest Rate: {avg_interest_rate:.2f}%"
        )
        self.chart_kpi_text.configure(state="normal")
        self.chart_kpi_text.delete("1.0", tk.END)
        self.chart_kpi_text.insert(tk.END, kpi)
        self.chart_kpi_text.configure(state="disabled")

    def clear_chart_filters(self) -> None:
        """Reset chart filters to defaults and refresh chart detail/KPI views."""
        self.chart_bank_entry.delete(0, tk.END)
        self.chart_loan_entry.delete(0, tk.END)
        self.chart_option_combo.set("Overview")
        self.chart_start_entry.delete(0, tk.END)
        self.chart_start_entry.insert(0, datetime.today().strftime(DATE_FORMAT))
        self.chart_end_entry.delete(0, tk.END)
        self.chart_end_entry.insert(0, (datetime.today() + timedelta(days=365)).strftime(DATE_FORMAT))
        self.chart_status_combo.set("All")
        self.refresh_chart_details()

    def display_chart(self) -> None:
        """Render and display the chart figure in a separate Matplotlib window."""
        if plt is None:
            raise ImportError("matplotlib is required for chart display. Install matplotlib to use charts.")
        rows = self._chart_filtered_rows()
        if not rows:
            raise ValueError("No chart data matches the selected filters.")
        filter_start = parse_date(self.chart_start_entry.get().strip(), "Chart Start Date")
        filter_end = parse_date(self.chart_end_entry.get().strip(), "Chart End Date")
        total_loan_amount = sum(float(row["loan_amount"]) for row in rows)
        total_remaining = sum(float(row["remaining_balance"]) for row in rows)
        selected_keys = {(str(row["bank_name"]).lower(), str(row["loan_number"]).lower()) for row in rows}

        status_counts = {"Paid": 0, "Upcoming": 0}
        for payment in self.data["payments"]:
            key = (str(payment.get("bank_name", "")).lower(), str(payment.get("loan_number", "")).lower())
            if key not in selected_keys:
                continue
            payment_status = str(payment.get("status", "Upcoming"))
            if payment_status not in status_counts:
                continue
            status_counts[payment_status] += 1

        figure = plt.figure(figsize=(13, 8))
        grid = figure.add_gridspec(2, 2, height_ratios=[1, 1.2])
        ax_pie = figure.add_subplot(grid[0, 0])
        pie_labels = ["Paid", "Upcoming"]
        pie_values = [status_counts[label] for label in pie_labels]
        ax_pie.pie(pie_values, labels=pie_labels, autopct="%1.1f%%", startangle=90)
        ax_pie.set_title("Loan Status Distribution")

        ax_amount_pie = figure.add_subplot(grid[0, 1])
        amount_labels = ["Total Loan Amount", "Remaining Balance"]
        amount_values = [total_loan_amount, total_remaining]
        ax_amount_pie.pie(amount_values, labels=amount_labels, autopct="%1.1f%%", startangle=90)
        ax_amount_pie.set_title("Loan Amount Metrics")

        ax_line = figure.add_subplot(grid[1, :])
        plotted = False
        for row in rows:
            loan_start = parse_date(str(row["start_date"]), "Start Date")
            loan_end = parse_date(str(row["end_date"]), "End Date")
            x_start = max(loan_start, filter_start)
            x_end = min(loan_end, filter_end)
            if x_end < x_start:
                continue
            ax_line.plot(
                [x_start, x_end],
                [float(row["loan_amount"]), float(row["remaining_balance"])],
                marker="o",
                label=f"{row['bank_name']} | {row['loan_number']}",
            )
            plotted = True
        if not plotted:
            raise ValueError("No line chart data available for the selected date range.")
        ax_line.set_title("Loan Amount Trend by Start/End Date")
        ax_line.set_xlabel("Date")
        ax_line.set_ylabel("Loan Amount")
        ax_line.grid(True, alpha=0.3)
        ax_line.legend(loc="upper right", fontsize=8)

        figure.tight_layout()
        plt.show(block=False)
        self.refresh_chart_details()

    def _write_table(self, file_path: str, headers: list[str], rows: list[list[object]]) -> None:
        """Write tabular data to CSV or XLSX based on file extension."""
        safe_rows = [self._sanitize_export_row(row) for row in rows]
        _, ext = os.path.splitext(file_path.lower())
        if ext == ".csv":
            with open(file_path, "w", newline="", encoding="utf-8") as handle:
                writer = csv.writer(handle)
                writer.writerow(headers)
                writer.writerows(safe_rows)
            return
        if ext == ".xlsx":
            if Workbook is None:
                raise ImportError("openpyxl is required for .xlsx export.")
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Export"
            sheet.append(headers)
            for row in safe_rows:
                sheet.append(row)
            workbook.save(file_path)
            return
        raise ValueError("Unsupported export extension. Use .csv or .xlsx.")

    def _sanitize_export_cell(self, value: object) -> object:
        """Neutralize spreadsheet formula prefixes in untrusted string cells before export."""
        if not isinstance(value, str):
            return value
        stripped = value.lstrip()
        if stripped.startswith(("=", "+", "-", "@")):
            return f"'{value}"
        return value

    def _sanitize_export_row(self, row: list[object]) -> list[object]:
        """Return a sanitized copy of an export row with spreadsheet-safe cell values."""
        return [self._sanitize_export_cell(value) for value in row]

    def _read_table(self, file_path: str) -> list[dict[str, object]]:
        """Read CSV or XLSX tabular data into a list of row dictionaries."""
        _, ext = os.path.splitext(file_path.lower())
        if ext == ".csv":
            with open(file_path, "r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle)
                return [dict(row) for row in reader]
        if ext == ".xlsx":
            if load_workbook is None:
                raise ImportError("openpyxl is required for .xlsx import.")
            workbook = load_workbook(file_path)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                return []
            headers = [str(value).strip() if value is not None else "" for value in rows[0]]
            parsed_rows: list[dict[str, object]] = []
            for raw in rows[1:]:
                mapped: dict[str, object] = {}
                for idx, header in enumerate(headers):
                    mapped[header] = raw[idx] if idx < len(raw) else None
                parsed_rows.append(mapped)
            return parsed_rows
        raise ValueError("Unsupported import extension. Use .csv or .xlsx.")


def main() -> None:
    """Create the Tk root, initialize the app, and start the main event loop."""
    root = tk.Tk()
    app = LoanScheduleApp(root)

    def _handle_tk_exception(exc_type, exc_value, exc_traceback) -> None:
        """Log uncaught Tk callback exceptions and show a sanitized user error dialog."""
        LOGGER.exception("Tk callback exception", exc_info=(exc_type, exc_value, exc_traceback))
        messagebox.showerror("Application Error", "An unexpected application error occurred.")

    root.report_callback_exception = _handle_tk_exception
    app.root.mainloop()


if __name__ == "__main__":
    main()
